"""Validation, schema migration, and idempotent metadata import."""

from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Iterable
from contextlib import contextmanager

from openpyxl import load_workbook


RIPENESS_NAMES = {"U": "unripe", "R": "ripe", "O": "overripe"}


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS banana_properties (
    banana_id TEXT PRIMARY KEY,
    arc_length_cm REAL NOT NULL,
    circumference_cm REAL NOT NULL,
    signal_path_cm REAL NOT NULL,
    initial_mass_g REAL NOT NULL,
    source_file TEXT NOT NULL,
    source_hash TEXT NOT NULL,
    imported_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS ripeness_observations (
    banana_id TEXT NOT NULL,
    measurement_date TEXT NOT NULL,
    study_day INTEGER NOT NULL,
    ripeness_code TEXT NOT NULL CHECK(ripeness_code IN ('U','R','O')),
    ripeness_class TEXT NOT NULL CHECK(ripeness_class IN ('unripe','ripe','overripe')),
    source_file TEXT NOT NULL,
    source_hash TEXT NOT NULL,
    imported_at TEXT NOT NULL,
    PRIMARY KEY (banana_id, measurement_date)
);
CREATE INDEX IF NOT EXISTS ripeness_date ON ripeness_observations(measurement_date);
CREATE TABLE IF NOT EXISTS metadata_imports (
    id INTEGER PRIMARY KEY,
    imported_at TEXT NOT NULL,
    dimensions_file TEXT NOT NULL,
    dimensions_hash TEXT NOT NULL,
    ripeness_file TEXT NOT NULL,
    ripeness_hash TEXT NOT NULL,
    study_start_date TEXT NOT NULL,
    property_rows INTEGER NOT NULL,
    ripeness_rows INTEGER NOT NULL,
    unmatched_ripeness_rows INTEGER NOT NULL
);

DROP VIEW IF EXISTS fit_parameters_long;
CREATE VIEW fit_parameters_long AS
SELECT
    f.id AS fit_id,
    c.id AS capture_id,
    c.banana_id,
    c.measurement_date,
    fc.display_name AS configuration,
    f.fruit_model,
    f.electrode_model,
    p.key AS parameter_name,
    CAST(p.value AS REAL) AS parameter_value
FROM fits AS f
JOIN captures AS c ON c.id = f.capture_id
JOIN fit_configs AS fc ON fc.id = f.config_id
JOIN json_each(f.parameters_json) AS p
WHERE f.status = 'complete';

DROP VIEW IF EXISTS fit_results_wide;
CREATE VIEW fit_results_wide AS
SELECT
    f.id AS fit_id,
    c.id AS capture_id,
    c.banana_id,
    c.measurement_date,
    fc.display_name AS configuration,
    f.fruit_model,
    f.electrode_model,
    f.nrmse,
    f.aicc,
    f.aic,
    f.bic,
    f.objective,
    json_extract(f.parameters_json, '$.r_inf') AS r_inf,
    json_extract(f.parameters_json, '$.delta_r') AS delta_r,
    json_extract(f.parameters_json, '$.tau_s') AS tau_s,
    json_extract(f.parameters_json, '$.alpha') AS alpha,
    json_extract(f.parameters_json, '$.r0') AS r0,
    json_extract(f.parameters_json, '$.r1') AS r1,
    json_extract(f.parameters_json, '$.r2') AS r2,
    json_extract(f.parameters_json, '$.c1_f') AS c1_f,
    json_extract(f.parameters_json, '$.c_e_f') AS c_e_f,
    json_extract(f.parameters_json, '$.sigma') AS sigma,
    json_extract(f.parameters_json, '$.r_e') AS r_e,
    json_extract(f.parameters_json, '$.q_e') AS q_e,
    json_extract(f.parameters_json, '$.n_e') AS n_e,
    f.parameters_json
FROM fits AS f
JOIN captures AS c ON c.id = f.capture_id
JOIN fit_configs AS fc ON fc.id = f.config_id
WHERE f.status = 'complete';

DROP VIEW IF EXISTS ml_dataset_view;
CREATE VIEW ml_dataset_view AS
SELECT
    fw.*,
    bp.arc_length_cm,
    bp.circumference_cm,
    bp.signal_path_cm,
    bp.initial_mass_g,
    ro.study_day,
    ro.ripeness_code,
    ro.ripeness_class
FROM fit_results_wide AS fw
LEFT JOIN banana_properties AS bp ON bp.banana_id = fw.banana_id
LEFT JOIN ripeness_observations AS ro
    ON ro.banana_id = fw.banana_id
   AND ro.measurement_date = fw.measurement_date;

DROP VIEW IF EXISTS ripeness_capture_status;
CREATE VIEW ripeness_capture_status AS
SELECT
    ro.*,
    CASE WHEN EXISTS (
        SELECT 1 FROM captures AS c
        WHERE c.is_current = 1
          AND c.banana_id = ro.banana_id
          AND c.measurement_date = ro.measurement_date
    ) THEN 1 ELSE 0 END AS has_eis_capture
FROM ripeness_observations AS ro;
"""


@dataclass(frozen=True)
class BananaProperty:
    banana_id: str
    arc_length_cm: float
    circumference_cm: float
    signal_path_cm: float
    initial_mass_g: float


@dataclass(frozen=True)
class RipenessObservation:
    banana_id: str
    measurement_date: str
    study_day: int
    ripeness_code: str

    @property
    def ripeness_class(self) -> str:
        return RIPENESS_NAMES[self.ripeness_code]


@dataclass(frozen=True)
class ImportPreview:
    properties: tuple[BananaProperty, ...]
    ripeness: tuple[RipenessObservation, ...]
    capture_dates: tuple[str, ...]
    unmatched_ripeness: tuple[RipenessObservation, ...]
    warnings: tuple[str, ...]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with _accessible_path(path) as readable:
        with readable.open("rb") as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(block)
    return digest.hexdigest()


@contextmanager
def _accessible_path(path: Path):
    """Stage OneDrive files when Python is denied direct access on Windows."""
    try:
        with path.open("rb") as stream:
            stream.read(1)
        yield path
        return
    except PermissionError:
        if os.name != "nt":
            raise
    with tempfile.TemporaryDirectory(prefix="eismetadata-", dir=Path.cwd()) as directory:
        staged = Path(directory) / path.name
        environment = os.environ.copy()
        environment["EISMETA_SOURCE"] = str(path)
        environment["EISMETA_DESTINATION"] = str(staged)
        completed = subprocess.run(
            [
                "powershell.exe", "-NoProfile", "-NonInteractive", "-Command",
                "Copy-Item -LiteralPath $env:EISMETA_SOURCE -Destination $env:EISMETA_DESTINATION -Force",
            ],
            capture_output=True,
            text=True,
            check=False,
            env=environment,
        )
        if completed.returncode != 0 or not staged.is_file():
            message = completed.stderr.strip() or "PowerShell could not stage the workbook"
            raise PermissionError(f"cannot read {path}: {message}")
        yield staged


def _banana_id(value: object) -> str:
    if value is None or str(value).strip() == "":
        raise ValueError("blank banana ID")
    text = str(value).strip().lower()
    if text.startswith("banana"):
        suffix = text[6:].strip()
    else:
        try:
            suffix = str(int(float(text)))
        except ValueError as exc:
            raise ValueError(f"invalid banana ID: {value!r}") from exc
    if not suffix:
        raise ValueError(f"invalid banana ID: {value!r}")
    return f"banana{suffix}"


def _number(value: object, field: str, banana_id: str) -> float:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{banana_id}: missing {field}")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{banana_id}: invalid {field}: {value!r}") from exc
    if number <= 0:
        raise ValueError(f"{banana_id}: {field} must be positive")
    return number


def read_dimensions(path: str | Path) -> tuple[BananaProperty, ...]:
    source = Path(path)
    with _accessible_path(source) as readable:
        workbook = load_workbook(readable, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    if not rows:
        raise ValueError("dimensions workbook is empty")
    headers = {str(value).strip().lower(): index for index, value in enumerate(rows[0]) if value is not None}
    expected = {
        "banana id": "banana_id",
        "arc length(cm)": "arc_length_cm",
        "circumference(cm)": "circumference_cm",
        "signal path(cm)": "signal_path_cm",
        "mass(g)": "initial_mass_g",
    }
    missing = [name for name in expected if name not in headers]
    if missing:
        raise ValueError("dimensions workbook is missing columns: " + ", ".join(missing))
    result: list[BananaProperty] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        banana_id = _banana_id(row[headers["banana id"]])
        if banana_id in seen:
            raise ValueError(f"duplicate dimensions row for {banana_id} at row {row_number}")
        seen.add(banana_id)
        result.append(BananaProperty(
            banana_id,
            _number(row[headers["arc length(cm)"]], "arc length", banana_id),
            _number(row[headers["circumference(cm)"]], "circumference", banana_id),
            _number(row[headers["signal path(cm)"]], "signal path", banana_id),
            _number(row[headers["mass(g)"]], "mass", banana_id),
        ))
    if not result:
        raise ValueError("dimensions workbook contains no data rows")
    return tuple(result)


def read_ripeness(path: str | Path, study_start: date) -> tuple[RipenessObservation, ...]:
    source = Path(path)
    with _accessible_path(source) as readable:
        workbook = load_workbook(readable, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    if len(rows) < 3:
        raise ValueError("ripeness workbook does not contain the expected two header rows")
    day_columns: list[tuple[int, int]] = []
    for column, value in enumerate(rows[1][1:], start=1):
        if value is None or str(value).strip() == "":
            continue
        try:
            day = int(float(value))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"invalid study day in column {column + 1}: {value!r}") from exc
        if day < 1:
            raise ValueError("study days must start at 1 or greater")
        day_columns.append((column, day))
    if not day_columns:
        raise ValueError("ripeness workbook has no study-day columns")
    observations: list[RipenessObservation] = []
    seen: set[tuple[str, int]] = set()
    for row_number, row in enumerate(rows[2:], start=3):
        if row[0] is None or str(row[0]).strip() == "":
            continue
        banana_id = _banana_id(row[0])
        for column, day in day_columns:
            value = row[column] if column < len(row) else None
            if value is None or str(value).strip() == "":
                continue
            code = str(value).strip().upper()
            if code not in RIPENESS_NAMES:
                raise ValueError(
                    f"ripeness row {row_number}, day {day}: expected U, R or O; found {value!r}"
                )
            key = (banana_id, day)
            if key in seen:
                raise ValueError(f"duplicate ripeness value for {banana_id}, day {day}")
            seen.add(key)
            observations.append(RipenessObservation(
                banana_id,
                (study_start + timedelta(days=day - 1)).isoformat(),
                day,
                code,
            ))
    if not observations:
        raise ValueError("ripeness workbook contains no observations")
    return tuple(observations)


def infer_study_start(database_path: str | Path) -> date:
    connection = sqlite3.connect(database_path)
    try:
        row = connection.execute(
            "SELECT MIN(measurement_date) FROM captures WHERE is_current=1"
        ).fetchone()
    finally:
        connection.close()
    if row is None or row[0] is None:
        raise ValueError("the database has no current captures from which to infer Day 1")
    return date.fromisoformat(str(row[0]))


def _capture_keys(connection: sqlite3.Connection) -> set[tuple[str, str]]:
    return {
        (str(row[0]), str(row[1]))
        for row in connection.execute(
            "SELECT banana_id,measurement_date FROM captures WHERE is_current=1"
        )
    }


def preview_metadata(database_path: str | Path, dimensions_path: str | Path,
                     ripeness_path: str | Path, study_start: date) -> ImportPreview:
    database = Path(database_path)
    if not database.is_file():
        raise ValueError(f"database not found: {database}")
    properties = read_dimensions(dimensions_path)
    ripeness = read_ripeness(ripeness_path, study_start)
    connection = sqlite3.connect(database)
    try:
        capture_keys = _capture_keys(connection)
        capture_dates = tuple(
            str(row[0]) for row in connection.execute(
                "SELECT DISTINCT measurement_date FROM captures WHERE is_current=1 ORDER BY measurement_date"
            )
        )
    finally:
        connection.close()
    unmatched = tuple(
        observation for observation in ripeness
        if (observation.banana_id, observation.measurement_date) not in capture_keys
    )
    property_ids = {item.banana_id for item in properties}
    capture_ids = {key[0] for key in capture_keys}
    warnings: list[str] = []
    missing_properties = sorted(capture_ids - property_ids)
    if missing_properties:
        warnings.append("Captures without dimensions: " + ", ".join(missing_properties))
    if unmatched:
        warnings.append(
            f"{len(unmatched)} ripeness observations have no EIS capture. "
            "They will be retained as metadata-only observations."
        )
    return ImportPreview(properties, ripeness, capture_dates, unmatched, tuple(warnings))


def import_metadata(database_path: str | Path, dimensions_path: str | Path,
                    ripeness_path: str | Path, study_start: date) -> ImportPreview:
    preview = preview_metadata(database_path, dimensions_path, ripeness_path, study_start)
    database = Path(database_path)
    dimensions = Path(dimensions_path).resolve()
    ripeness_file = Path(ripeness_path).resolve()
    dimensions_hash = _hash_file(dimensions)
    ripeness_hash = _hash_file(ripeness_file)
    imported_at = _now()
    connection = sqlite3.connect(database)
    try:
        connection.execute("BEGIN IMMEDIATE")
        connection.executescript(SCHEMA)
        connection.executemany(
            """INSERT INTO banana_properties(
                   banana_id,arc_length_cm,circumference_cm,signal_path_cm,initial_mass_g,
                   source_file,source_hash,imported_at
               ) VALUES(?,?,?,?,?,?,?,?)
               ON CONFLICT(banana_id) DO UPDATE SET
                   arc_length_cm=excluded.arc_length_cm,
                   circumference_cm=excluded.circumference_cm,
                   signal_path_cm=excluded.signal_path_cm,
                   initial_mass_g=excluded.initial_mass_g,
                   source_file=excluded.source_file,
                   source_hash=excluded.source_hash,
                   imported_at=excluded.imported_at""",
            [(
                item.banana_id, item.arc_length_cm, item.circumference_cm,
                item.signal_path_cm, item.initial_mass_g, str(dimensions),
                dimensions_hash, imported_at,
            ) for item in preview.properties],
        )
        connection.executemany(
            """INSERT INTO ripeness_observations(
                   banana_id,measurement_date,study_day,ripeness_code,ripeness_class,
                   source_file,source_hash,imported_at
               ) VALUES(?,?,?,?,?,?,?,?)
               ON CONFLICT(banana_id,measurement_date) DO UPDATE SET
                   study_day=excluded.study_day,
                   ripeness_code=excluded.ripeness_code,
                   ripeness_class=excluded.ripeness_class,
                   source_file=excluded.source_file,
                   source_hash=excluded.source_hash,
                   imported_at=excluded.imported_at""",
            [(
                item.banana_id, item.measurement_date, item.study_day,
                item.ripeness_code, item.ripeness_class, str(ripeness_file),
                ripeness_hash, imported_at,
            ) for item in preview.ripeness],
        )
        connection.execute(
            """INSERT INTO metadata_imports(
                   imported_at,dimensions_file,dimensions_hash,ripeness_file,ripeness_hash,
                   study_start_date,property_rows,ripeness_rows,unmatched_ripeness_rows
               ) VALUES(?,?,?,?,?,?,?,?,?)""",
            (
                imported_at, str(dimensions), dimensions_hash, str(ripeness_file), ripeness_hash,
                study_start.isoformat(), len(preview.properties), len(preview.ripeness),
                len(preview.unmatched_ripeness),
            ),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return preview
